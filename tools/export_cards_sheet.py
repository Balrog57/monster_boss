"""Exporte src/frontend/godot/data/cards.json en classeur Excel : une feuille par type de carte, une ligne par carte.
Usage : python tools/export_cards_sheet.py [sortie.xlsx]   (defaut : docs/cartes.xlsx ; necessite openpyxl)"""
import json
import os
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from cards_io import ROOT, load_cards

SHEETS = [("bosses", "Boss"), ("rooms", "Salles"), ("spells", "Sorts"), ("heroes", "Héros"),
          ("items", "Objets"), ("minibosses", "Minibosses")]
FIRST = ["id", "name", "subtitle", "set", "quantity", "implemented"]


def cell(v):
    if isinstance(v, bool):
        return "oui" if v else "non"
    if isinstance(v, list):
        if all(isinstance(x, dict) and "level" in x for x in v):
            return "\n".join(f"Niv. {x['level']} : {x.get('description', '')}" for x in v)
        return ", ".join(str(x) for x in v)
    if isinstance(v, dict):
        return json.dumps(v, ensure_ascii=False)
    return v


def write_sheet(ws, cards):
    seen = Counter(k for c in cards for k in c)
    cols = [k for k in FIRST if k in seen] + sorted(k for k in seen if k not in FIRST)
    ws.append(cols)
    for c in cards:
        ws.append([cell(c.get(k)) for k in cols])
    for cell_ in ws[1]:
        cell_.font = Font(bold=True, color="FFFFFF")
        cell_.fill = PatternFill("solid", fgColor="4F2A2A")
    for i, k in enumerate(cols, 1):
        width = max(len(str(r[i - 1].value or "")) for r in ws.iter_rows(min_row=1))
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 8), 70)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions


def main():
    data = load_cards()
    out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, "docs", "cartes.xlsx")
    wb = Workbook()
    summary = wb.active
    summary.title = "Résumé"
    summary.append(["Type", "Cartes distinctes", "Exemplaires (quantity)", "Implémentées"])
    for key, title in SHEETS:
        cards = data[key]
        summary.append([title, len(cards), sum(c.get("quantity", 1) for c in cards),
                        sum(1 for c in cards if c.get("implemented"))])
        write_sheet(wb.create_sheet(title), cards)
    summary.append(["Total", *(f"=SUM({get_column_letter(i)}2:{get_column_letter(i)}{len(SHEETS) + 1})" for i in (2, 3, 4))])
    for c in summary[1]:
        c.font = Font(bold=True)
    for c in summary[summary.max_row]:
        c.font = Font(bold=True)
    for i in range(1, 5):
        summary.column_dimensions[get_column_letter(i)].width = 24
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
